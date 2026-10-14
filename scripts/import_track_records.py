"""
Script import dữ liệu từ file "Track Record.xlsx" vào Supabase
Đọc sheet "Clean Data", parse từng dòng và insert vào bảng track_records
"""

import sys
import json
import urllib.request
import urllib.parse
import openpyxl
from datetime import datetime, date

# ─── CONFIG ───────────────────────────────────────────────────────────────────
SUPABASE_URL  = "https://etoetjxqvxjobohpdpny.supabase.co"
SUPABASE_KEY  = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImV0b2V0anhxdnhqb2JvaHBkcG55Iiwicm9sZSI6ImFub24iLCJpYXQiOjE3NTkzNzU5OTcsImV4cCI6MjA3NDk1MTk5N30.ALvQwFMd4WnM3TQ3rATk_uZ3u_l4PuJkhONkXo8OpcI"
TABLE         = "track_records"
EXCEL_PATH    = r"d:\ANTIGRAVITY APPs\Web ROV Workshop\Track Record.xlsx"
SHEET_NAME    = "Clean Data"
BATCH_SIZE    = 50   # số dòng insert mỗi lần gọi API

# ─── HELPERS ──────────────────────────────────────────────────────────────────
def safe_str(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None

def safe_float(v) -> float | None:
    if v is None:
        return None
    try:
        f = float(str(v).replace(",", ""))
        return f
    except Exception:
        return None

def safe_int(v) -> int | None:
    f = safe_float(v)
    return int(f) if f is not None else None

def safe_date(v) -> str | None:
    """Chuyển giá trị ô ngày tháng sang ISO string yyyy-mm-dd"""
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if not s or s.startswith("="):   # công thức chưa tính
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None

def supabase_insert(rows: list[dict]) -> dict:
    """Gọi Supabase REST API để insert nhiều dòng"""
    url = f"{SUPABASE_URL}/rest/v1/{TABLE}"
    payload = json.dumps(rows, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=payload,
        method="POST",
        headers={
            "apikey":        SUPABASE_KEY,
            "Authorization": f"Bearer {SUPABASE_KEY}",
            "Content-Type":  "application/json",
            "Prefer":        "return=minimal",
        }
    )
    try:
        with urllib.request.urlopen(req) as resp:
            return {"ok": True, "status": resp.status}
    except urllib.error.HTTPError as e:
        body = e.read().decode("utf-8", errors="replace")
        return {"ok": False, "status": e.code, "body": body}

# ─── MAIN ─────────────────────────────────────────────────────────────────────
def main():
    print(f"[INFO] Đọc file Excel: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[SHEET_NAME]

    # Đọc header (hàng 1)
    headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
    print(f"[INFO] Sheet '{SHEET_NAME}': {ws.max_row - 1} dòng dữ liệu, {ws.max_column} cột")
    print(f"[INFO] Headers: {headers}")

    rows_to_insert = []
    skipped = 0
    total = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        # Bỏ qua dòng hoàn toàn rỗng
        if all(v is None for v in row):
            skipped += 1
            continue

        # Map cột theo thứ tự trong file Excel:
        # NO. | CLIENT | CONTRACT | PROJECT TASK | START DATE | END DATE |
        # DURATION(DAYS) | CLIENT CONTACT PERSON | CLIENT CONTACT PHONE |
        # PROJECT SOW | LOCATION | TOTAL PROJECT COST (USD) |
        # WATER DEPTH MIN | WATER DEPTH MAX | ROV SYSTEM | VESSEL/RIG | CLASS
        def get(idx):
            return row[idx] if idx < len(row) else None

        record = {
            "no":                    safe_int(get(0)),
            "client":                safe_str(get(1)),
            "contract":              safe_str(get(2)),
            "project_task":          safe_str(get(3)),
            "start_date":            safe_date(get(4)),
            "end_date":              safe_date(get(5)),
            # cot 6 la DURATION - duoc tinh tu dong boi DB, bo qua
            "client_contact_person": safe_str(get(7)),
            "client_contact_phone":  safe_str(get(8)),
            "project_sow":           safe_str(get(9)),
            "location":              safe_str(get(10)),
            "total_project_cost_usd":safe_float(get(11)),
            "water_depth_min_msw":   safe_float(get(12)),
            "water_depth_max_msw":   safe_float(get(13)),
            "rov_system":            safe_str(get(14)),
            "vessel_rig":            safe_str(get(15)),
            "class":                 safe_str(get(16)),
            "created_by":            "Import from Excel",
        }

        # KEEP None values - Supabase requires all rows in a batch to have identical keys
        rows_to_insert.append(record)
        total += 1

    print(f"[INFO] Total valid rows: {total} | Skipped: {skipped}")
    print(f"[INFO] Deleting old data before re-import...")

    # Delete all existing rows first to avoid duplicates
    delete_url = f"{SUPABASE_URL}/rest/v1/{TABLE}?id=gte.0"
    delete_req = urllib.request.Request(
        delete_url,
        method="DELETE",
        headers={
            "apikey":        SUPABASE_KEY,
            "Authorization": f"Bearer {SUPABASE_KEY}",
        }
    )
    try:
        with urllib.request.urlopen(delete_req) as resp:
            print(f"[INFO] Old data deleted (status {resp.status})")
    except urllib.error.HTTPError as e:
        print(f"[WARN] Could not delete old data: {e.read().decode()}")

    print(f"[INFO] Starting batch insert (batch size={BATCH_SIZE})...")

    success = 0
    fail = 0
    for i in range(0, len(rows_to_insert), BATCH_SIZE):
        batch = rows_to_insert[i:i + BATCH_SIZE]
        result = supabase_insert(batch)
        batch_end = min(i + BATCH_SIZE, total)
        if result["ok"]:
            success += len(batch)
            print(f"  [OK ] Rows {i+1}–{batch_end} → inserted ({len(batch)} rows)")
        else:
            fail += len(batch)
            print(f"  [ERR] Rows {i+1}–{batch_end} → {result.get('status')} {result.get('body', '')[:200]}")

    print(f"\n[DONE] Thành công: {success}/{total} dòng | Thất bại: {fail} dòng")

if __name__ == "__main__":
    main()
